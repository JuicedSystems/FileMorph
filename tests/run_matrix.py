"""Convert every sample to every format FileMorph offers for it and validate the output.

usage: python tests/run_matrix.py <samples_dir> <out_dir> [--only pdf,docx] [--jobs 2]
Set FILEMORPH_NO_LIBREOFFICE=1 to test the built-in (no Office) engines.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
import warnings
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
warnings.filterwarnings("ignore")


def validate(p: Path, fmt: str) -> str | None:
    """Return an error string, or None if the file looks like a valid `fmt`."""
    if p.is_dir():
        return None
    if p.stat().st_size == 0:
        return "empty file"
    try:
        if fmt == "trace":
            import re as _re
            text = p.read_text(encoding="utf-8")
            fills = set(_re.findall(r'fill="([^"]+)"', text))
            if "<path" not in text or fills - {"#000000"}:
                return f"trace: expected black paths only, got fills {fills}"
            fmt = "svg"
        if fmt in ("pdf", "svg"):
            import pymupdf
            with pymupdf.open(p) as d:
                if d.page_count < 1:
                    return "no pages"
        elif fmt in ("png", "jpg", "webp", "gif", "tiff", "bmp", "ico", "avif"):
            from PIL import Image
            with Image.open(p) as im:
                im.load()
                if im.width < 8:
                    return f"tiny image {im.size}"
        elif fmt == "docx":
            import docx
            d = docx.Document(p)
            if not any(par.text.strip() for par in d.paragraphs) and not d.tables and not d.inline_shapes:
                return "docx has no content"
        elif fmt == "pptx":
            from pptx import Presentation
            if len(Presentation(p).slides) < 1:
                return "no slides"
        elif fmt in ("xlsx",):
            import openpyxl
            wb = openpyxl.load_workbook(p)
            if not any(ws.max_row >= 1 and ws.cell(1, 1).value is not None for ws in wb.worksheets):
                return "empty workbook"
        elif fmt in ("csv", "tsv", "txt", "md", "html", "json", "rtf"):
            text = p.read_text(encoding="utf-8-sig", errors="replace")
            if len(text.strip()) < 3:
                return "almost empty text"
            if fmt == "json":
                json.loads(text)
        elif fmt in ("doc", "xls", "ppt"):
            if p.read_bytes()[:4] != b"\xd0\xcf\x11\xe0":
                return "not an OLE2 office file"
        elif fmt in ("odt", "ods", "odp", "epub"):
            import zipfile
            if not zipfile.is_zipfile(p):
                return "not a zip container"
        elif fmt in ("mp3", "wav", "m4a", "flac", "ogg", "opus", "mp4", "mov", "mkv", "webm", "avi"):
            from filemorph.core.converters.media import probe
            info = probe(p)
            if info["duration"] < 0.3:
                return f"duration {info['duration']}"
            if fmt in ("mp4", "mov", "mkv", "webm", "avi") and not info["video"]:
                return "no video stream"
    except Exception as e:  # noqa: BLE001
        return f"invalid: {type(e).__name__}: {e}"
    return None


def one(sample: str, dst: str, out_root: str):
    from filemorph.core.engine import convert_file
    from filemorph.core.formats import detect
    src = Path(sample)
    out = Path(out_root) / f"{src.suffix[1:]}_to_{dst}"
    out.mkdir(parents=True, exist_ok=True)
    t = time.time()
    try:
        files = convert_file(src, dst, out, {})
        errs = [e for f in files if (e := validate(f, dst if f.suffix[1:].lower() in ("md",) or True else dst))]
        status = "FAIL" if errs else "ok"
        detail = "; ".join(errs) if errs else f"{len(files)} file(s)"
    except Exception as e:  # noqa: BLE001
        status, detail = "ERR", f"{type(e).__name__}: {str(e)[:160]}"
    return detect(src).key, dst, status, detail, time.time() - t


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("samples")
    ap.add_argument("out")
    ap.add_argument("--only", default="")
    ap.add_argument("--jobs", type=int, default=2)
    a = ap.parse_args()
    from filemorph.core.converters import build_registry
    from filemorph.core.formats import detect
    reg = build_registry()
    only = set(filter(None, a.only.split(",")))
    tasks = []
    for s in sorted(Path(a.samples).iterdir()):
        f = detect(s)
        if f is None or (only and f.key not in only):
            continue
        for d in sorted(reg.targets(f.key)):
            tasks.append((str(s), d))
    print(f"{len(tasks)} conversions", flush=True)
    results = []
    with ProcessPoolExecutor(a.jobs) as ex:
        futs = [ex.submit(one, s, d, a.out) for s, d in tasks]
        for fu in as_completed(futs):
            r = fu.result()
            results.append(r)
            if r[2] != "ok":
                print(f"  {r[2]:4} {r[0]:>5} -> {r[1]:<5} {r[3]}", flush=True)
    ok = sum(r[2] == "ok" for r in results)
    slow = sorted(results, key=lambda r: -r[4])[:5]
    print(f"\n{ok}/{len(results)} passed. Slowest: " + ", ".join(f"{r[0]}->{r[1]} {r[4]:.1f}s" for r in slow))
    Path(a.out, "results.json").write_text(json.dumps(results, indent=1))
    sys.exit(0 if ok == len(results) else 1)


if __name__ == "__main__":
    main()
