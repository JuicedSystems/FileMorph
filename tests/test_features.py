"""Focused checks for OCR, SHRINK, recipes and watch folders.

usage: python tests/test_features.py [name-filter]
Plain asserts, no test framework needed. Everything happens in a temp folder; recipes and the
watch ledger go to a temp FILEMORPH_DATA, never the user's real ones.
"""
from __future__ import annotations

import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import traceback
import warnings
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))
warnings.filterwarnings("ignore")
TMP = Path(tempfile.mkdtemp(prefix="fm-features-"))
os.environ["FILEMORPH_DATA"] = str(TMP / "data")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from filemorph.core.engine import convert_file  # noqa: E402
from filemorph.core.registry import ConversionError  # noqa: E402

TESTS = []


def test(fn):
    TESTS.append(fn)
    return fn


def fresh(name: str) -> Path:
    d = TMP / name
    shutil.rmtree(d, ignore_errors=True)
    d.mkdir(parents=True)
    return d


def scan_pdf(d: Path) -> Path:
    import pymupdf
    from make_samples import make_scan
    buf = io.BytesIO()
    make_scan().convert("L").save(buf, "JPEG", quality=80)
    doc = pymupdf.open()
    doc.new_page(width=612, height=792).insert_image(pymupdf.Rect(0, 0, 612, 792), stream=buf.getvalue())
    doc.save(d / "scan.pdf")
    return d / "scan.pdf"


def photo(size=(3000, 2000)):
    """A noisy, detailed picture: big as JPEG, like a phone photo."""
    import numpy as np
    from PIL import Image, ImageFilter
    rng = np.random.default_rng(1)
    a = rng.integers(0, 255, (size[1] // 8, size[0] // 8, 3), dtype=np.uint8)
    im = Image.fromarray(a).resize(size, Image.BICUBIC).filter(ImageFilter.GaussianBlur(1.2))
    noise = Image.fromarray(rng.integers(0, 40, (size[1], size[0], 3), dtype=np.uint8))
    return Image.blend(im, noise, 0.25)


# ------------------------------------------------------------------ OCR
@test
def ocr_scanned_pdf_to_text_word_excel():
    d = fresh("ocr")
    src = scan_pdf(d)
    txt = convert_file(src, "txt", d / "o")[0].read_text(encoding="utf-8")
    for w in ("INVOICE 10442", "Northwind Traders", "Widget", "Gadget", "149.00", "Thank you"):
        assert w in txt, (w, txt)
    import openpyxl
    ws = openpyxl.load_workbook(convert_file(src, "xlsx", d / "o")[0]).worksheets[0]
    rows = [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
    assert ("Item", "Qty", "Price") in rows, rows
    assert ("Widget", 2, 14.5) in rows and ("Gadget", 10, 120) in rows, rows
    import docx
    doc = docx.Document(convert_file(src, "docx", d / "o")[0])
    assert any("INVOICE 10442" in p.text for p in doc.paragraphs)
    cells = [[c.text for c in r.cells] for t in doc.tables for r in t.rows]
    assert ["Item", "Qty", "Price"] in cells and ["Total", "", "149.00"] in cells, cells


@test
def ocr_searchable_pdf_keeps_the_scan():
    import pymupdf
    d = fresh("ocr2")
    out = convert_file(scan_pdf(d), "ocrpdf", d / "o")[0]
    assert out.name == "scan - searchable.pdf", out.name
    with pymupdf.open(out) as doc:
        page = doc[0]
        assert "Northwind" in page.get_text() and len(page.get_images()) == 1
        # the invisible words sit on top of the picture's words (selection lines up)
        w = next(w for w in page.get_text("words") if w[4] == "Northwind")
        assert 90 < w[0] < 120 and 90 < w[1] < 112, w  # page points: 290 px, 280 px at 200 dpi
    try:
        convert_file(out, "ocrpdf", d / "o2")
        raise AssertionError("an already searchable PDF should be refused")
    except ConversionError as e:
        assert "already has selectable text" in str(e)


@test
def ocr_photo_of_document_to_text():
    from make_samples import make_scan
    d = fresh("ocr3")
    make_scan().save(d / "page.jpg", quality=85)
    text = convert_file(d / "page.jpg", "txt", d / "o")[0].read_text(encoding="utf-8")
    assert "Northwind" in text and "149.00" in text, text


@test
def ocr_never_runs_on_digital_pdfs():
    import pymupdf
    from filemorph.core.converters import ocr
    d = fresh("ocr4")
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Plain digital text, nothing to read.")
    doc.save(d / "digital.pdf")
    real = ocr.ocr_page
    ocr.ocr_page = lambda page: (_ for _ in ()).throw(AssertionError("OCR ran on a digital PDF"))
    try:
        assert "Plain digital" in convert_file(d / "digital.pdf", "txt", d / "o")[0].read_text(encoding="utf-8")
        convert_file(d / "digital.pdf", "xlsx", d / "o")
    finally:
        ocr.ocr_page = real
    assert not ocr.looks_scanned(d / "digital.pdf")
    assert ocr.looks_scanned(scan_pdf(d))


@test
def ocr_blank_scan_gives_a_friendly_error():
    import pymupdf
    from PIL import Image
    d = fresh("ocr5")
    buf = io.BytesIO()
    Image.new("RGB", (1200, 1600), "white").save(buf, "JPEG")
    doc = pymupdf.open()
    doc.new_page().insert_image(pymupdf.Rect(0, 0, 612, 792), stream=buf.getvalue())
    doc.save(d / "blank.pdf")
    try:
        convert_file(d / "blank.pdf", "txt", d / "o")
        raise AssertionError("expected an error")
    except ConversionError as e:
        assert "No text found" in str(e), e
    assert not (d / "o").exists() or not any((d / "o").iterdir())


# ------------------------------------------------------------------ SHRINK
def _smaller(src: Path, out: Path, ratio=0.9):
    assert out.exists(), out
    assert out.stat().st_size < src.stat().st_size * ratio, (out.stat().st_size, src.stat().st_size)


@test
def shrink_photos_keep_format_and_extension():
    from PIL import Image
    d = fresh("shrink_img")
    im = photo()
    im.save(d / "Holiday.jpeg", quality=97)
    out = convert_file(d / "Holiday.jpeg", "shrink", d / "o")[0]
    assert out.name == "Holiday - small.jpeg", out.name
    _smaller(d / "Holiday.jpeg", out, 0.6)
    with Image.open(out) as o:
        assert o.format == "JPEG" and max(o.size) == 2560, (o.format, o.size)
    out = convert_file(d / "Holiday.jpeg", "shrink", d / "o", {"shrink_strength": "strong"})[0]
    with Image.open(out) as o:
        assert max(o.size) == 1600
    im.save(d / "pic.webp", quality=95)
    with Image.open(convert_file(d / "pic.webp", "shrink", d / "o")[0]) as o:
        assert o.format == "WEBP"
    # a screenshot-like PNG becomes a palette PNG
    from PIL import ImageDraw
    shot = Image.new("RGB", (1600, 1000), "#f3f4f6")
    dr = ImageDraw.Draw(shot)
    for i in range(40):
        dr.rectangle((20 + i * 38, 40, 50 + i * 38, 900), fill=(i * 6, 90, 200 - i * 4))
        dr.text((30 + i * 38, 920), str(i), fill="black")
    shot.save(d / "shot.png")
    out = convert_file(d / "shot.png", "shrink", d / "o")[0]
    _smaller(d / "shot.png", out, 0.95)
    with Image.open(out) as o:
        assert o.format == "PNG"


@test
def shrink_pdf_with_big_pictures():
    import pymupdf
    d = fresh("shrink_pdf")
    buf = io.BytesIO()
    photo().save(buf, "PNG")
    doc = pymupdf.open()
    for i in range(2):
        pg = doc.new_page()
        pg.insert_image(pymupdf.Rect(36, 36, 576, 400), stream=buf.getvalue())
        pg.insert_text((72, 450), f"Page {i + 1} text stays text.")
    doc.save(d / "brochure.pdf")
    out = convert_file(d / "brochure.pdf", "shrink", d / "o")[0]
    assert out.name == "brochure - small.pdf"
    _smaller(d / "brochure.pdf", out, 0.5)
    with pymupdf.open(out) as o:
        assert o.page_count == 2 and "stays text" in o[1].get_text()


@test
def shrink_office_files_pictures():
    import docx
    from docx.shared import Inches
    from pptx import Presentation
    from pptx.util import Inches as PI
    d = fresh("shrink_office")
    photo().save(d / "big.png")
    doc = docx.Document()
    doc.add_paragraph("Report with a huge picture")
    doc.add_picture(str(d / "big.png"), width=Inches(5))
    doc.save(d / "report.docx")
    out = convert_file(d / "report.docx", "shrink", d / "o")[0]
    _smaller(d / "report.docx", out, 0.6)
    assert docx.Document(out).paragraphs[0].text == "Report with a huge picture"
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(d / "big.png"), PI(0), PI(0), PI(10))
    prs.save(d / "deck.pptx")
    out = convert_file(d / "deck.pptx", "shrink", d / "o")[0]
    _smaller(d / "deck.pptx", out, 0.6)
    assert len(Presentation(out).slides) == 1


@test
def shrink_video_and_audio():
    from filemorph.core.converters.media import ffmpeg_exe, probe
    d = fresh("shrink_media")
    ff = ffmpeg_exe()
    subprocess.run([ff, "-y", "-loglevel", "error", "-f", "lavfi", "-i", "testsrc2=duration=4:size=1920x1080:rate=30",
                    "-f", "lavfi", "-i", "sine=duration=4", "-shortest", "-c:v", "libx264", "-crf", "10",
                    "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "256k", str(d / "clip.mp4")], check=True)
    out = convert_file(d / "clip.mp4", "shrink", d / "o")[0]
    _smaller(d / "clip.mp4", out, 0.5)
    info = probe(out)
    assert info["video"] and info["audio"] and abs(info["duration"] - 4) < 0.5, info
    subprocess.run([ff, "-y", "-loglevel", "error", "-f", "lavfi", "-i", "anoisesrc=d=5", "-c:a", "libmp3lame",
                    "-b:a", "320k", str(d / "talk.mp3")], check=True)
    out = convert_file(d / "talk.mp3", "shrink", d / "o")[0]
    _smaller(d / "talk.mp3", out, 0.6)


@test
def shrink_refuses_when_already_compact():
    d = fresh("shrink_compact")
    photo((400, 300)).save(d / "tiny.jpg", quality=45, optimize=True)  # small and already lossy
    try:
        convert_file(d / "tiny.jpg", "shrink", d / "o")
        raise AssertionError("expected 'already as small'")
    except ConversionError as e:
        assert str(e).startswith("Already as small"), e
    assert not (d / "o").exists() or not any((d / "o").iterdir())


@test
def shrink_is_never_chained():
    from filemorph.core.converters import build_registry
    reg = build_registry()
    for k in ("heic", "bmp", "gif", "wav", "flac", "avi", "csv", "doc", "svg"):
        assert "shrink" not in reg.targets(k), k  # would need a format change: not SHRINK
    for k in ("pdf", "jpg", "png", "docx", "pptx", "xlsx", "mp4", "mp3"):
        assert len(reg.plan(k, "shrink")) == 1, k


# ------------------------------------------------------------------ recipes core
@test
def recipe_names_and_storage():
    from filemorph.core import recipes as rc
    src = Path("C:/x/Holiday photo.heic")
    t0 = time.mktime((2026, 10, 23, 9, 5, 7, 0, 0, -1))
    assert rc.render_name("", src) is None
    assert rc.render_name("{name} {date}", src, now=t0) == "Holiday photo 2026-10-23"
    assert rc.render_name("{name}_{n}_{time}", src, 3, now=t0) == "Holiday photo_3_09-05-07"
    assert rc.render_name('a:b/c*?"<>|.', src) == "a_b_c______"
    assert rc.render_name("   ...  ", src) is None
    a = rc.Recipe(name="Phone → JPG", target="jpg", options={"quality": 85}, rename="{name}")
    b = rc.Recipe(name="Scans", target="ocrpdf", watch_dir=str(TMP), watching=True)
    rc.save([a, b])
    back = rc.load()
    assert [r.name for r in back] == ["Phone → JPG", "Scans"] and back[0].options == {"quality": 85}
    assert back[1].watching and rc.find(back, a.id).rename == "{name}"
    assert "watching" in back[1].summary() and "→ OCR" in back[1].summary()
    rc.path().write_text("{not json", encoding="utf-8")
    assert rc.load() == [] and rc.path().with_suffix(".json.bad").exists()
    rc.path().write_text(json.dumps({"recipes": [{"name": "bad", "target": "nope"}, {"name": 3},
                                                  {"name": "ok", "target": "png", "options": {"x": [1]}}]}),
                         encoding="utf-8")
    good = rc.load()
    assert [r.name for r in good] == ["ok"] and good[0].options == {}
    rc.save([])


@test
def convert_file_uses_recipe_names():
    from PIL import Image
    d = fresh("rename")
    Image.new("RGB", (64, 64), "red").save(d / "a.png")
    out = convert_file(d / "a.png", "jpg", d / "o", out_stem="Trip 2026-10-23")[0]
    assert out.name == "Trip 2026-10-23.jpg"
    again = convert_file(d / "a.png", "jpg", d / "o", out_stem="Trip 2026-10-23")[0]
    assert again.name == "Trip 2026-10-23 (1).jpg"  # never overwrites


@test
def watch_rules_ledger_and_settling():
    from PIL import Image
    from filemorph.core import recipes as rc
    from filemorph.core.converters import build_registry
    reg = build_registry()
    d = fresh("watch_rules")
    Image.new("RGB", (32, 32)).save(d / "old.png")
    r = rc.Recipe(name="to jpg", target="jpg", watch_dir=str(d), watching=True)
    led = rc.Ledger(TMP / "data" / "ledger-test.json")
    assert led.baseline(r) == 1
    for n in ("new.png", "~$lock.png", ".hidden.png", "dl.png.crdownload", "notes.xyz", "photo.jpg"):
        (d / n).write_bytes((d / "old.png").read_bytes())
    (d / "notes.xyz").write_text("not a picture", encoding="utf-8")  # unknown type: ignored
    os.utime(d / "new.png", (time.time() - 10 * 86400,) * 2)  # copied in with an old date: still new
    names = [p.name for p, _ in rc.candidates(r, reg, led)]
    assert names == ["new.png"], names
    led.add(r.id, d / "new.png", [d / "new.jpg"])
    assert rc.candidates(r, reg, led) == []
    (d / "new.png").write_bytes(b"x" + (d / "old.png").read_bytes())  # edited: counts as new again
    assert [p.name for p, _ in rc.candidates(r, reg, led)] == ["new.png"]
    led.save()
    assert rc.Ledger(led.file).seen(r.id, d / "new.jpg")
    s = rc.Settle(quiet=1.0)
    st = (d / "old.png").stat()
    assert not s.ready(d / "old.png", st, now=100.0)
    assert not s.ready(d / "old.png", st, now=100.5)
    assert s.ready(d / "old.png", st, now=101.2)
    shr = rc.Recipe(name="smaller", target="shrink", watch_dir=str(d), watching=True)
    (d / "pic - small.jpg").write_bytes(b"x")
    assert "pic - small.jpg" not in [p.name for p, _ in rc.candidates(shr, reg, led)]


@test
def inbox_carries_recipes():
    from filemorph import instance
    old = os.environ.get("LOCALAPPDATA")
    os.environ["LOCALAPPDATA"] = str(fresh("inbox"))
    try:
        instance.post(["a.png"])
        instance.post(["b.png"], "abcd1234")
        instance.post(["c.png"])
        instance.post(["d.png"], "abcd1234")
        batches, seen = instance.drain()
        assert seen and batches == [(["a.png", "c.png"], None), (["b.png", "d.png"], "abcd1234")], batches
        instance.post([])
        assert instance.drain() == ([([], None)], True)
    finally:
        if old is not None:
            os.environ["LOCALAPPDATA"] = old


@test
def explorer_recipe_menu_registry():
    if sys.platform != "win32":
        return
    import winreg
    from filemorph.core import recipes as rc
    from filemorph.core import shell
    base = r"Software\FileMorphTest\shell\FileMorphRecipes"  # never the live context menu
    exe = r"C:\Apps\FileMorph.exe"
    real = shell._RECIPE_KEYS
    shell._RECIPE_KEYS = [base]
    try:
        a = rc.Recipe(name="Scans & receipts", target="ocrpdf")
        shell.install_recipes([a, rc.Recipe(name="Smaller", target="shrink")], exe=exe)
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, base) as k:
            assert winreg.QueryValueEx(k, "MUIVerb")[0] == "FileMorph recipes"
            assert winreg.QueryValueEx(k, "SubCommands")[0] == ""
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, base + r"\shell\r00") as k:
            assert winreg.QueryValueEx(k, "MUIVerb")[0] == "Scans && receipts"
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, base + r"\shell\r01\command") as k:
            assert winreg.QueryValueEx(k, "")[0].startswith(f'"{exe}" --recipe ')
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, base + r"\shell\r00\command") as k:
            assert winreg.QueryValueEx(k, "")[0] == f'"{exe}" --recipe {a.id} "%1"'
        shell.install_recipes([], exe=exe)  # no recipes: the menu goes away
        try:
            winreg.OpenKey(winreg.HKEY_CURRENT_USER, base).Close()
            raise AssertionError("menu left behind")
        except FileNotFoundError:
            pass
    finally:
        shell._RECIPE_KEYS = real
        shell._delete_tree(winreg, winreg.HKEY_CURRENT_USER, r"Software\FileMorphTest")


# ------------------------------------------------------------------ Qt: watch service & controller
def _pump(seconds: float, until=None):
    from PySide6.QtCore import QCoreApplication
    end = time.time() + seconds
    while time.time() < end:
        QCoreApplication.processEvents()
        if until and until():
            return True
        time.sleep(0.02)
    return bool(until and until())


def _app():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


@test
def watch_service_converts_new_files_once():
    from PIL import Image
    from filemorph.core import recipes as rc
    from filemorph.core.converters import build_registry
    from filemorph.ui.node.watch import WatchService
    _app()
    d = fresh("watch_live")
    Image.new("RGB", (40, 40), "blue").save(d / "already.png")
    r = rc.Recipe(name="Shrink here", target="shrink", watch_dir=str(d), watching=True,
                  options={"shrink_strength": "balanced"})
    recipes = [r]
    svc = WatchService(build_registry(), lambda: recipes, ledger=rc.Ledger(TMP / "data" / "ledger-live.json"))
    svc.ledger.baseline(r)
    done = []
    svc.converted.connect(lambda rr, src, outs, err: done.append((Path(src).name, [Path(o).name for o in outs], err)))
    bursts = []
    svc.burst_done.connect(lambda name, ok, bad: bursts.append((name, ok, bad)))
    svc.refresh()
    photo((2400, 1600)).save(d / "incoming.jpg", quality=97)  # arrives after watching started
    assert _pump(30, lambda: done), "the new file was never converted"
    _pump(8, lambda: bursts)
    assert done == [("incoming.jpg", ["incoming - small.jpg"], "")], done
    assert bursts == [("Shrink here", 1, 0)], bursts
    _pump(6)  # its output sits in the watched folder too: it must not be shrunk again
    assert len(done) == 1 and sorted(p.name for p in d.iterdir()) == ["already.png", "incoming - small.jpg",
                                                                      "incoming.jpg"]
    svc.set_paused(True)
    photo((800, 600)).save(d / "paused.jpg", quality=97)
    _pump(5)
    assert len(done) == 1
    svc.set_paused(False)
    assert _pump(20, lambda: len(done) == 2), "resumed watching should catch up"
    assert svc.stop()


@test
def controller_recipes_scans_and_shrink_caption():
    from PIL import Image
    from PySide6.QtCore import QSettings
    from filemorph.core import recipes as rc
    app = _app()
    s = QSettings("FileMorph", "FileMorph-tests")
    s.clear()
    from filemorph.ui.node import controller as ctl
    d = fresh("controller")
    out_dir = fresh("controller_out")
    rc.save([rc.Recipe(name="PNG to web JPG", target="jpg", options={"quality": 80}, save_dir=str(out_dir),
                       rename="{name} web", sources=["png"])])
    c = ctl.MorphController(app)
    c.settings = s
    try:
        Image.new("RGB", (300, 200), "green").save(d / "logo.png")
        c.add_files([d / "logo.png"])
        _pump(0.3)
        chips = [ch.key for ch in c.node.chips]
        assert chips[0].startswith("recipe:") and c.node.chips[0].label == "JPG" and c.node.chips[0].star, chips
        assert c.node.chips[0].hint == "PNG to web JPG"
        c.node.chip_picked.emit(chips[0])
        assert _pump(30, lambda: c.node.mode != "converting" and c.worker is None)
        assert [Path(o).name for o in c.last_outputs] == ["logo web.jpg"], c.last_outputs
        assert Path(c.last_outputs[0]).parent == out_dir
        # the recipe chip only appears for files it can convert
        (d / "notes.txt").write_text("hello", encoding="utf-8")
        c.node.show_idle()
        c.add_files([d / "notes.txt"])
        _pump(0.2)
        assert not any(ch.key.startswith("recipe:") for ch in c.node.chips)
        # a scanned PDF puts OCR first (the photo recipe isn't offered for a PDF)
        c.node.show_idle()
        c.add_files([scan_pdf(d)])
        _pump(0.2)
        keys = [ch.key for ch in c.node.chips]
        assert keys[0] == "ocrpdf", keys
        # SHRINK shows before -> after
        c.node.show_idle()
        photo((2400, 1600)).save(d / "big.jpg", quality=97)
        c._out_dir = lambda: out_dir
        c.add_files([d / "big.jpg"])
        _pump(0.2)
        assert "shrink" in [ch.key for ch in c.node.chips]
        c.node.chip_picked.emit("shrink")
        assert _pump(40, lambda: c.worker is None and c.node.mode in ("done", "error"))
        cap = [t for k, t in c.node.caption.parts]
        assert c.node.mode == "done" and any("MB →" in t or "KB →" in t for t in cap), cap
        # recipes from Explorer arrive through receive_paths
        r = rc.load()[0]
        c.node.show_idle()
        c.receive_paths([d / "logo.png"], r.id)
        assert _pump(30, lambda: c.worker is None and any(p.name == "logo web (1).jpg" for p in out_dir.iterdir()))
        # panel: recipes list, editor, save
        c.show_recipes()
        _pump(0.5)
        assert c.panel.stack.currentWidget() is c.panel.recipes_pg
        c.panel.show_recipe_editor(None, c.panel._editor_ctx, {"target": "shrink", "options": {}, "name": "Smaller"})
        _pump(0.3)
        assert c.panel.ed_target.currentData() == "shrink" and "shrink_strength" in c.panel.ed_opt_widgets
        c.panel.ed_rename.setText("{name} (small)")
        new = c.panel.editor_values()
        assert new.name == "Smaller" and new.options == {"shrink_strength": "balanced"} and new.rename
        c._save_recipe(new)
        assert [x.name for x in rc.load()] == ["PNG to web JPG", "Smaller"]
        c._delete_recipe(new.id)
        assert [x.name for x in rc.load()] == ["PNG to web JPG"]
    finally:
        c.watch.stop()
        c.node.hide()
        c.panel.hide()
        c.tray.hide()
        s.clear()
        rc.save([])


def main():
    flt = sys.argv[1] if len(sys.argv) > 1 else ""
    fails = 0
    for fn in TESTS:
        if flt and flt not in fn.__name__:
            continue
        t = time.time()
        try:
            fn()
            print(f"  ok    {fn.__name__}  ({time.time() - t:.1f}s)", flush=True)
        except Exception:  # noqa: BLE001
            fails += 1
            print(f"  FAIL  {fn.__name__}\n" + "".join("        " + ln for ln in
                                                       traceback.format_exc().splitlines(True)), flush=True)
    shutil.rmtree(TMP, ignore_errors=True)
    print(f"\n{'all passed' if not fails else f'{fails} failed'}")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
