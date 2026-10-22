"""Spreadsheet & tabular data conversions."""
from __future__ import annotations

import csv
import html as htmllib
import json
import re
from pathlib import Path

from ..registry import ConversionError, Ctx, Registry
from .documents import read_text, render_html_pdf

SHEET_IN = ["xlsx", "xls", "ods", "csv", "tsv", "json"]
SHEET_OUT = ["xlsx", "ods", "csv", "tsv", "json", "html"]


def read_sheets(src: Path, fmt: str) -> dict:
    """Return {sheet_name: DataFrame}."""
    import pandas as pd
    try:
        if fmt in ("xlsx", "xls", "ods"):
            engine = {"xlsx": "openpyxl", "xls": "xlrd", "ods": "odf"}[fmt]
            sheets = pd.read_excel(src, sheet_name=None, engine=engine, header=None)
            sheets = {k: _promote_header(v) for k, v in sheets.items()}
        elif fmt in ("csv", "tsv"):
            text = read_text(src)
            if fmt == "tsv":
                delim = "\t"
            else:
                try:
                    delim = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
                except csv.Error:
                    delim = ","
            import io
            sheets = {src.stem[:31] or "Sheet1": pd.read_csv(io.StringIO(text), sep=delim, dtype=str,
                                                             keep_default_na=False, engine="python")}
            # ';'-separated files come from locales that write decimals as 1,5
            sheets = {k: _infer_numbers(v, decimal_comma=delim == ";") for k, v in sheets.items()}
        elif fmt == "json":
            data = json.loads(read_text(src))
            sheets = _json_to_frames(data, src.stem)
        else:
            raise ConversionError(f"Can't read {fmt} as a table.")
    except ConversionError:
        raise
    except Exception as e:
        raise ConversionError("This file couldn't be read as a table — it may be damaged or not tabular.") from e
    sheets = {k: v for k, v in sheets.items() if not (v.empty and len(v.columns) == 0)} or sheets
    if not sheets:
        raise ConversionError("The file is empty.")
    return sheets


def _promote_header(df):
    """First non-empty row becomes the header (read with header=None to keep blank top rows sane)."""
    df = df.dropna(how="all").dropna(axis=1, how="all")
    if df.empty:
        return df
    header = [str(c) if c == c and c is not None else f"Column {i + 1}" for i, c in enumerate(df.iloc[0])]
    seen: dict[str, int] = {}
    for i, h in enumerate(header):  # de-duplicate
        if h in seen:
            seen[h] += 1
            header[i] = f"{h} ({seen[h]})"
        else:
            seen[h] = 0
    out = df.iloc[1:].reset_index(drop=True)
    out.columns = header
    return out


_PLAIN = re.compile(r"^-?\d+(\.\d+)?([eE][-+]?\d+)?$")
_THOUSANDS = re.compile(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$")
_EU_PLAIN = re.compile(r"^-?\d+(,\d+)?$")
_EU_THOUSANDS = re.compile(r"^-?\d{1,3}(\.\d{3})+(,\d+)?$")
_LEADING_ZERO = re.compile(r"^-?0\d")


def _normalise_number(v: str, decimal_comma: bool) -> str | None:
    """'1,234.5' -> '1234.5' (or '1.234,5' -> '1234.5' with decimal commas); None if not a number."""
    v = v.strip()
    if decimal_comma:
        if _EU_PLAIN.match(v) or _EU_THOUSANDS.match(v):
            return v.replace(".", "").replace(",", ".")
        return None
    if _PLAIN.match(v):
        return v
    if _THOUSANDS.match(v):
        return v.replace(",", "")
    return None


def _infer_numbers(df, decimal_comma: bool = False):
    """Turn a text column into numbers only when that loses nothing: every value must be an
    unambiguous number, and columns with leading zeros (ZIP codes, IDs) or more digits than a
    double can hold stay text."""
    import pandas as pd
    for c in df.columns:
        col = df[c].astype(str)
        vals = [v.strip() for v in col if v.strip() != ""]
        if not vals:
            continue
        if any(_LEADING_ZERO.match(v) for v in vals):
            continue
        if any(sum(ch.isdigit() for ch in v.lstrip("-0")) > 15 for v in vals):
            continue
        norm = [_normalise_number(v, decimal_comma) for v in vals]
        if any(n is None for n in norm):
            continue
        numeric = pd.to_numeric(col.map(lambda v: _normalise_number(v, decimal_comma) if v.strip() else None),
                                errors="coerce")
        if all(("." not in n and "e" not in n.lower()) for n in norm):
            numeric = numeric.astype("Int64")  # keep 7 as 7, not 7.0
        df[c] = numeric
    return df


def _json_to_frames(data, stem: str) -> dict:
    import pandas as pd
    if isinstance(data, dict):
        lists = {k: v for k, v in data.items() if isinstance(v, list)}
        if lists and len(lists) == len(data):  # {"sheet": [rows], ...}
            return {str(k)[:31]: pd.json_normalize(v) if v and isinstance(v[0], dict) else pd.DataFrame({k: v})
                    for k, v in lists.items()}
        return {stem[:31] or "Sheet1": pd.json_normalize(data)}
    if isinstance(data, list):
        if data and all(isinstance(r, dict) for r in data):
            return {stem[:31] or "Sheet1": pd.json_normalize(data)}
        if data and all(isinstance(r, list) for r in data):
            return {stem[:31] or "Sheet1": _promote_header(pd.DataFrame(data))}
        return {stem[:31] or "Sheet1": pd.DataFrame({"value": data})}
    raise ConversionError("This JSON isn't shaped like a table (expected a list of records).")


def _safe_sheet(name: str) -> str:
    for ch in '[]:*?/\\':
        name = name.replace(ch, "_")
    return name[:31] or "Sheet"


def _unique_sheet_names(names) -> list[str]:
    """Excel-safe, case-insensitively unique: 'a/b' and 'a_b' become 'a_b' and 'a_b (2)'."""
    out, seen = [], set()
    for n in names:
        base = _safe_sheet(str(n))
        cand, i = base, 2
        while cand.lower() in seen:
            suffix = f" ({i})"
            cand = base[: 31 - len(suffix)] + suffix
            i += 1
        seen.add(cand.lower())
        out.append(cand)
    return out


def _multi_paths(dst: Path, names: list[str]) -> list[Path]:
    if len(names) == 1:
        return [dst]
    return [dst.with_name(f"{dst.stem} - {n}{dst.suffix}") for n in _unique_sheet_names(names)]


def _tidy_dates(df):
    """Date-only datetime columns as 2024-01-05 (not 2024-01-05 00:00:00) in text outputs."""
    import pandas as pd
    out = df
    for c in df.columns:
        col = df[c]
        if pd.api.types.is_datetime64_any_dtype(col):
            nn = col.dropna()
            if len(nn) and (nn.dt.normalize() == nn).all():
                if out is df:
                    out = df.copy()
                out[c] = col.dt.strftime("%Y-%m-%d").where(col.notna(), "")
    return out


def sheets_to_html(sheets: dict, title: str, for_pdf: bool = False) -> str:
    parts = []
    for name, df in sheets.items():
        if len(sheets) > 1:
            parts.append(f"<h2>{htmllib.escape(str(name))}</h2>")
        parts.append(_tidy_dates(df).to_html(index=False, na_rep="", border=0, escape=True))
    css = "" if for_pdf else (
        "body{font-family:'Segoe UI',system-ui,sans-serif;margin:32px;color:#1f2328}"
        "table{border-collapse:collapse;font-size:14px;margin-bottom:32px}"
        "th,td{border:1px solid #d0d7de;padding:6px 10px;text-align:left}"
        "th{background:#f6f8fa;position:sticky;top:0}tr:nth-child(even) td{background:#fbfbfc}")
    return (f"<!doctype html><html><head><meta charset='utf-8'><title>{htmllib.escape(title)}</title>"
            f"<style>{css}</style></head><body>{''.join(parts)}</body></html>")


def sheet_convert(src: Path, dst: Path, ctx: Ctx):
    sheets = read_sheets(src, ctx.src_fmt)
    ctx.progress(0.4, "Writing data")
    d = ctx.dst_fmt
    if d == "xlsx":
        _write_xlsx(sheets, dst)
    elif d == "ods":
        import pandas as pd
        with pd.ExcelWriter(dst, engine="odf") as w:
            for sn, df in zip(_unique_sheet_names(sheets), sheets.values()):
                df.to_excel(w, sheet_name=sn, index=False)
    elif d in ("csv", "tsv"):
        outs = _multi_paths(dst, [str(k) for k in sheets])
        for (name, df), out in zip(sheets.items(), outs):
            # utf-8-sig so Excel opens accented characters correctly
            _tidy_dates(df).to_csv(out, index=False, sep="," if d == "csv" else "\t", encoding="utf-8-sig")
        return outs
    elif d == "json":
        def records(df):
            return json.loads(df.to_json(orient="records", date_format="iso", force_ascii=False))
        data = records(next(iter(sheets.values()))) if len(sheets) == 1 else {str(k): records(v) for k, v in sheets.items()}
        dst.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    elif d == "html":
        dst.write_text(sheets_to_html(sheets, src.stem), encoding="utf-8")
    else:  # pragma: no cover
        raise ConversionError(f"Unsupported target {d}")


def sheet_to_pdf(src: Path, dst: Path, ctx: Ctx):
    sheets = read_sheets(src, ctx.src_fmt)
    wide = max(len(df.columns) for df in sheets.values()) > 7
    opts = dict(ctx.options)
    if wide:
        opts["page_size"] = "letter-l" if opts.get("page_size", "letter").startswith("letter") else "a4-l"
    ctx2 = Ctx(ctx.src_fmt, ctx.dst_fmt, ctx.tmpdir, opts, ctx.progress_cb, ctx.cancel_event)
    render_html_pdf(sheets_to_html(sheets, src.stem, for_pdf=True), dst, ctx2)


def _write_xlsx(sheets: dict, dst: Path):
    import pandas as pd
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(dst, engine="openpyxl") as w:
        for sn, df in zip(_unique_sheet_names(sheets), sheets.values()):
            df.to_excel(w, sheet_name=sn, index=False)
            ws = w.sheets[sn]
            # Text that starts with "=" must stay text, never become a live formula (CSV injection)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if cell.data_type == "f":
                        cell.value = str(cell.value)
                        cell.data_type = "s"
                        cell.quotePrefix = True
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="EEF0F4")
            ws.freeze_panes = "A2"
            for i, col in enumerate(df.columns, 1):
                sample = [str(col)] + [str(v) for v in df.iloc[:200, i - 1].tolist()]
                ws.column_dimensions[get_column_letter(i)].width = min(60, max(9, max(map(len, sample)) + 2))


def register(reg: Registry) -> None:
    for s in SHEET_IN:
        for d in SHEET_OUT:
            if s != d:
                reg.add(s, d, sheet_convert, multi=d in ("csv", "tsv"), cost=1.0, label="Converting table")
    reg.add(SHEET_IN, "pdf", sheet_to_pdf, cost=1.3, label="Laying out PDF")
